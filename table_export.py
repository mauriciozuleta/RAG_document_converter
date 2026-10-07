"""Export recognized PDF tables without inferring numeric cell types."""
import csv
from html.parser import HTMLParser
from pathlib import Path


class TableParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.tables = []
        self.rows = None
        self.row = None
        self.cell = None
        self.spans = []

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == 'table':
            if self.rows is not None:
                raise ValueError('Nested tables require manual review.')
            self.rows, self.spans = [], []
        elif tag == 'tr' and self.rows is not None:
            self.row = []
        elif tag in ('td', 'th') and self.row is not None:
            self.cell = []
            self.span = tuple(max(1, int(attrs.get(key, 1))) for key in ('rowspan', 'colspan'))
        elif tag == 'br' and self.cell is not None:
            self.cell.append('\n')

    def handle_data(self, data):
        if self.cell is not None:
            self.cell.append(data)

    def handle_endtag(self, tag):
        if tag in ('td', 'th') and self.cell is not None:
            self.row.append((''.join(self.cell).strip(), self.span))
            self.cell = None
        elif tag == 'tr' and self.row is not None:
            self.rows.append(self.row)
            self.row = None
        elif tag == 'table' and self.rows is not None:
            grid, occupied, merges = [], {}, []
            for r, row in enumerate(self.rows):
                c = 0
                for value, (height, width) in row:
                    while (r, c) in occupied:
                        c += 1
                    if height > 10000 or width > 1000:
                        raise ValueError('Unreasonable table span; review the OCR result.')
                    for rr in range(r, r + height):
                        for cc in range(c, c + width):
                            if (rr, cc) in occupied:
                                raise ValueError('Overlapping table cells; review the OCR result.')
                            occupied[rr, cc] = value if (rr, cc) == (r, c) else ''
                    if height > 1 or width > 1:
                        merges.append((r + 1, c + 1, r + height, c + width))
                    c += width
            if occupied:
                height = max(r for r, c in occupied) + 1
                width = max(c for r, c in occupied) + 1
                grid = [[occupied.get((r, c), '') for c in range(width)] for r in range(height)]
                self.tables.append((grid, merges))
            self.rows = None


def export_tables(data, output_dir, base_name, fmt):
    tables = []
    for section in data['sections']:
        parser = TableParser()
        parser.feed(section.get('ocr_markdown', ''))
        if parser.rows is not None:
            raise ValueError('Incomplete recognized table; review the OCR result.')
        detected = [(t['rows'], t.get('merges', [])) for t in section.get('tables', [])] + parser.tables
        for number, (rows, merges) in enumerate(detected, 1):
            tables.append((f"page-{section['source_page']}-table-{number}", rows, merges))
    if not tables:
        raise ValueError('No structured tables detected on the selected PDF pages. Review the OCR checkpoints or try different pages.')
    destination = Path(output_dir)
    destination.mkdir(parents=True, exist_ok=True)
    paths = []
    if fmt == 'csv':
        for name, rows, merges in tables:
            path = destination / f'{base_name}_{name}.csv'
            with path.open('w', encoding='utf-8-sig', newline='') as stream:
                csv.writer(stream).writerows(rows)
            paths.append(str(path))
    elif fmt == 'xlsx':
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.remove(workbook.active)
        for name, rows, merges in tables:
            sheet = workbook.create_sheet(name)
            for r, row in enumerate(rows, 1):
                for c, value in enumerate(row, 1):
                    cell = sheet.cell(r, c, value)
                    cell.data_type = 's'
                    cell.number_format = '@'
            for r1, c1, r2, c2 in merges:
                sheet.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
            sheet.freeze_panes = 'A2'
        path = destination / f'{base_name}.xlsx'
        workbook.save(path)
        paths.append(str(path))
    else:
        raise ValueError('Table format must be csv or xlsx.')
    print(f'[INFO] Exported {len(tables)} tables. Review recognized cells against the PDF. CSV merged cells use blank continuation cells; import CSV columns as text to preserve identifiers.')
    for path in paths:
        print(f'[INFO] Output written: {path}')
    return paths
